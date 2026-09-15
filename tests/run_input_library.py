#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
GovEnt AgentShield - Input Guard 样例库批量测试
================================================

作用：
1. 读取 data/ 下的 Excel 攻击样例库；
2. 从“样例库”工作表逐行读取正式测试样例；
3. 自动调用 src/input_guard.py 中的 InputGuard；
4. 自动生成每条样例的历史 JSON；
5. 自动生成批量测试记录 CSV；
6. 自动统计 TP/TN/FP/FN、风险等级准确率和动作准确率；
7. 不修改原始样例库 Excel。

依赖：
    pip install openpyxl
"""

from __future__ import annotations

import argparse
import csv
import json
import sys
from collections import Counter
from pathlib import Path
from typing import Any, Iterable


# =========================
# 1. 定位项目目录并导入 InputGuard
# =========================

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SRC_DIR = PROJECT_ROOT / "src"

if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

try:
    from input_guard import InputGuard
except ImportError as exc:
    raise SystemExit(
        "无法导入 src/input_guard.py。\n"
        f"请确认文件存在：{SRC_DIR / 'input_guard.py'}\n"
        f"原始错误：{exc}"
    ) from exc


# =========================
# 2. 样例库字段
# =========================

REQUIRED_COLUMNS = (
    "sample_id",
    "scenario",
    "input_text",
    "attack_type",
    "expected_level",
    "expected_action",
    "input_source",
    "attack_target",
    "expected_reason",
    "reference_source",
)

OUTPUT_COLUMNS = (
    # 样例库原始字段
    "sample_id",
    "scenario",
    "module",
    "ground_truth",
    "input_source",
    "input_or_target",
    "expected_attack_type",
    "attack_target",
    "expected_risk_level",
    "expected_action",
    "expected_reason",
    "reference_source",

    # 实际运行信息
    "event_id",
    "run_id",
    "test_date",
    "timestamp",
    "tester",

    # Input Guard 实际结果
    "risk_score",
    "risk_level",
    "actual_action",
    "actual_risk_type",
    "matched_rules",
    "matched_categories",
    "reason",

    # 对比结果
    "classification_result",
    "detection_is_correct",
    "risk_level_is_correct",
    "action_is_correct",
    "error_type",

    # 证据与版本
    "json_file_name",
    "screenshot_file_name",
    "evidence_status",
    "model_version",
    "rule_version",
    "dataset_version",
    "template_version",
)


def _text(value: Any) -> str:
    """Excel 单元格安全转字符串。"""
    if value is None:
        return ""
    return str(value).strip()


def _nullable_choice(value: Any) -> str | None:
    """空单元格转换为 None。"""
    text = _text(value)
    return text if text else None


def _csv_value(value: Any) -> str | int | float:
    """把列表/字典转换成适合 CSV 的文本。"""
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    if isinstance(value, (str, int, float)):
        return value
    return str(value)


# =========================
# 3. 读取 Excel 样例库
# =========================

def load_sample_library(
    workbook_path: Path,
    sheet_name: str = "样例库",
) -> list[dict[str, Any]]:
    """
    读取第二周最新 Excel 攻击样例库。

    第1行：标题
    第2行：字段名
    第3行开始：正式样例
    """

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "缺少 openpyxl，无法读取 .xlsx 样例库。\n"
            "请执行：python -m pip install openpyxl"
        ) from exc

    workbook_path = workbook_path.resolve()

    if not workbook_path.exists():
        raise FileNotFoundError(
            f"找不到样例库：{workbook_path}"
        )

    wb = load_workbook(
        workbook_path,
        read_only=True,
        data_only=True,
    )

    if sheet_name not in wb.sheetnames:
        raise ValueError(
            f"样例库中不存在工作表“{sheet_name}”。"
            f"当前工作表：{', '.join(wb.sheetnames)}"
        )

    ws = wb[sheet_name]

    # 第2行为字段名
    header_values = [
        _text(cell.value)
        for cell in ws[2]
    ]

    header_map = {
        name: index
        for index, name in enumerate(header_values)
        if name
    }

    missing = [
        column
        for column in REQUIRED_COLUMNS
        if column not in header_map
    ]

    if missing:
        raise ValueError(
            "样例库缺少必要字段："
            + ", ".join(missing)
        )

    samples: list[dict[str, Any]] = []

    # 第3行开始，共60条正式样例
    for row in ws.iter_rows(
        min_row=3,
        values_only=True
    ):

        sample_id = _text(
            row[header_map["sample_id"]]
        )

        if not sample_id:
            continue

        attack_type = _text(
            row[header_map["attack_type"]]
        )

        # 最新样例库 → Input Guard 内部统一格式
        ground_truth = (
            "normal"
            if attack_type == "normal_request"
            else "attack"
        )

        record = {
            "sample_id": sample_id,

            "scenario": row[
                header_map["scenario"]
            ],

            # 固定属于输入检测模块
            "module": "input_guard",

            # normal_request = 正常，其余均视为攻击
            "ground_truth": ground_truth,

            "input_source": row[
                header_map["input_source"]
            ],

            # 新字段 input_text
            # 转成旧脚本内部字段 input_or_target
            "input_or_target": row[
                header_map["input_text"]
            ],

            # 新字段 attack_type
            # 转成旧脚本内部字段
            "attack_type_or_operation": attack_type,

            "attack_target": row[
                header_map["attack_target"]
            ],

            # expected_level → expected_risk_level
            "expected_risk_level": row[
                header_map["expected_level"]
            ],

            "expected_action": row[
                header_map["expected_action"]
            ],

            "expected_reason": row[
                header_map["expected_reason"]
            ],

            "reference_source": row[
                header_map["reference_source"]
            ],
        }

        samples.append(record)

    wb.close()

    return samples


# =========================
# 4. 单条样例运行
# =========================

def run_one_sample(
    guard: InputGuard,
    sample: dict[str, Any],
    *,
    run_id: str,
    tester: str,
) -> tuple[dict[str, Any], dict[str, Any]]:
    """运行一条样例并生成检测结果和测试记录行。"""

    sample_id = _text(sample["sample_id"])
    scenario = _text(sample["scenario"]) or "政企通用场景"
    ground_truth = _nullable_choice(sample["ground_truth"])
    expected_risk_level = _nullable_choice(sample["expected_risk_level"])
    expected_action = _nullable_choice(sample["expected_action"])
    input_text = _text(sample["input_or_target"])

    if not input_text:
        raise ValueError(f"{sample_id} 的 input_or_target 为空。")

    result = guard.detect(
        input_text,
        sample_id=sample_id,
        run_id=run_id,
        tester=tester,
        scenario=scenario,
        ground_truth=ground_truth,
        expected_risk_level=expected_risk_level,
        expected_action=expected_action,
        is_valid_test="有效",
    )

    # 额外判断“风险等级是否与样例库预期完全一致”
    risk_level_is_correct: str | None
    if expected_risk_level is None:
        risk_level_is_correct = None
    else:
        risk_level_is_correct = (
            "true"
            if result["risk_level"] == expected_risk_level
            else "false"
        )

    # 保存最新结果
    guard.save_result(result)

    # 保存历史 JSON（Input Guard 独立历史目录）
    history_dir = guard.output_dir / "input_guard" / "history"
    history_path = history_dir / f"{result['event_id']}.json"
    guard.save_result(result, history_path)

    record = {
        "sample_id": sample_id,
        "scenario": scenario,
        "module": "input_guard",
        "ground_truth": ground_truth,
        "input_source": _text(sample["input_source"]),
        "input_or_target": input_text,
        "expected_attack_type": _text(
            sample["attack_type_or_operation"]
        ),
        "attack_target": _text(sample["attack_target"]),
        "expected_risk_level": expected_risk_level,
        "expected_action": expected_action,
        "expected_reason": _text(sample["expected_reason"]),
        "reference_source": _text(sample["reference_source"]),

        "event_id": result["event_id"],
        "run_id": result["run_id"],
        "test_date": result["test_date"],
        "timestamp": result["timestamp"],
        "tester": result["tester"],

        "risk_score": result["risk_score"],
        "risk_level": result["risk_level"],
        "actual_action": result["actual_action"],
        "actual_risk_type": result["risk_type"],
        "matched_rules": result["matched_rules"],
        "matched_categories": result["matched_categories"],
        "reason": result["reason"],

        "classification_result": result["classification_result"],
        "detection_is_correct": result["detection_is_correct"],
        "risk_level_is_correct": risk_level_is_correct,
        "action_is_correct": result["action_is_correct"],
        "error_type": result["error_type"],

        "json_file_name": history_path.name,
        "screenshot_file_name": result["screenshot_file_name"],
        "evidence_status": result["evidence_status"],
        "model_version": result["model_version"],
        "rule_version": result["rule_version"],
        "dataset_version": result["dataset_version"],
        "template_version": result["template_version"],
    }

    return result, record


# =========================
# 5. 保存批量测试记录
# =========================

def save_batch_records(
    records: Iterable[dict[str, Any]],
    output_path: Path,
) -> Path:
    """将本轮样例库测试记录保存为 CSV。"""
    output_path = output_path.resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)

    with output_path.open(
        "w",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        writer = csv.DictWriter(
            file,
            fieldnames=list(OUTPUT_COLUMNS),
            extrasaction="ignore",
        )
        writer.writeheader()

        for record in records:
            writer.writerow({
                column: _csv_value(record.get(column))
                for column in OUTPUT_COLUMNS
            })

    return output_path


# =========================
# 6. 统计汇总
# =========================

def print_batch_summary(
    records: list[dict[str, Any]],
    output_path: Path,
) -> None:
    """打印本轮样例库测试统计。"""
    total = len(records)

    confusion = Counter(
        _text(record.get("classification_result"))
        for record in records
        if _text(record.get("classification_result"))
    )

    detection_correct = sum(
        1
        for record in records
        if _text(record.get("detection_is_correct")).lower() == "true"
    )

    risk_correct = sum(
        1
        for record in records
        if _text(record.get("risk_level_is_correct")).lower() == "true"
    )

    action_correct = sum(
        1
        for record in records
        if _text(record.get("action_is_correct")).lower() == "true"
    )

    print("\n" + "=" * 72)
    print("GovEnt AgentShield - 样例库批量测试完成")
    print("=" * 72)
    print(f"样例总数：{total}")

    if total:
        print(
            f"检测正确：{detection_correct}/{total} "
            f"({detection_correct / total:.2%})"
        )
        print(
            f"风险等级完全一致：{risk_correct}/{total} "
            f"({risk_correct / total:.2%})"
        )
        print(
            f"处置动作完全一致：{action_correct}/{total} "
            f"({action_correct / total:.2%})"
        )

    print(
        "混淆矩阵："
        f"TP={confusion.get('TP', 0)}，"
        f"TN={confusion.get('TN', 0)}，"
        f"FP={confusion.get('FP', 0)}，"
        f"FN={confusion.get('FN', 0)}"
    )

    wrong_samples = [
        record["sample_id"]
        for record in records
        if (
            _text(record.get("risk_level_is_correct")).lower() == "false"
            or _text(record.get("action_is_correct")).lower() == "false"
        )
    ]

    if wrong_samples:
        print(
            "需要重点复核的样例："
            + ", ".join(wrong_samples)
        )
    else:
        print("所有样例的风险等级和处置动作均与预期一致。")

    print(f"测试记录：{output_path}")
    print("=" * 72)


# =========================
# 7. 命令行入口
# =========================

def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="批量读取 Excel 攻击样例库并测试 Input Guard"
    )
    parser.add_argument(
        "--library",
        type=Path,
        required=True,
        help="攻击样例库 .xlsx 文件路径",
    )
    parser.add_argument(
        "--sheet",
        default="样例库",
        help="样例库工作表名称",
    )
    parser.add_argument(
        "--run-id",
        default="R02",
        help="本轮测试编号",
    )
    parser.add_argument(
        "--tester",
        default="赵安意",
        help="实际执行正式复测的人员",
    )
    parser.add_argument(
        "--output",
        type=Path,
        help=(
            "批量测试记录 CSV；省略时保存到 "
            "output/statistics/input_library_test_records.csv"
        ),
    )
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    try:
        guard = InputGuard(project_root=PROJECT_ROOT)

        samples = load_sample_library(
            args.library,
            sheet_name=args.sheet,
        )

        if not samples:
            print("样例库中没有找到 input_guard 样例。")
            return 1

        print(
            f"已读取 {len(samples)} 条 input_guard 样例，"
            "开始正式批量检测……"
        )

        records: list[dict[str, Any]] = []

        for index, sample in enumerate(samples, start=1):
            sample_id = _text(sample["sample_id"])
            try:
                result, record = run_one_sample(
                    guard,
                    sample,
                    run_id=args.run_id,
                    tester=args.tester,
                )
                records.append(record)

                print(
                    f"[{index:02d}/{len(samples):02d}] "
                    f"{sample_id} | "
                    f"{result['risk_level']} | "
                    f"{result['actual_action']} | "
                    f"{result['classification_result'] or '-'}"
                )

            except Exception as exc:
                print(
                    f"[ERROR] {sample_id} 测试失败：{exc}",
                    file=sys.stderr,
                )

        output_path = (
            args.output
            if args.output is not None
            else PROJECT_ROOT
            / "output"
            / "statistics"
            / "input_library_test_records.csv"
        )

        output_path = save_batch_records(
            records,
            output_path,
        )

        print_batch_summary(
            records,
            output_path,
        )

        return 0

    except (
        FileNotFoundError,
        ValueError,
        RuntimeError,
        OSError,
    ) as exc:
        print(f"运行失败：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
